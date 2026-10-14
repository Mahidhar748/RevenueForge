"""Generates two fictional company packs (the files a company uploads during onboarding).

Meridian Integration Labs  - integration & healthcare interoperability specialist (strong fit for an integration RFP)
Cobalt Ridge Technologies  - cloud, data & BFSI digital company (partial fit for a fraud-intelligence RFP)

The two packs deliberately use different spreadsheet layouts (list vs. rating grid, day vs. hourly
rates) to exercise RevenueForge's column mapping. All companies, people and projects are fictional.

Run:  python -m app.seed.generate_packs   (writes to ../sample-data)
"""
from __future__ import annotations

import random
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[3] / "sample-data"


def _xlsx(path: Path, sheet: str, header: list[str], rows: list[list], widths: list[int] | None = None, title: str | None = None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    r0 = 1
    if title:
        ws.cell(1, 1, title).font = Font(bold=True, size=13)
        r0 = 3
    for j, h in enumerate(header, start=1):
        c = ws.cell(r0, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(rows, start=r0 + 1):
        for j, v in enumerate(row, start=1):
            c = ws.cell(i, j, v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for j, w in enumerate(widths or [22] * len(header), start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ============================================================================ Meridian
MERIDIAN = {
    "name": "Meridian Integration Labs",
    "slug": "meridian",
    "industry": "IT Services – Integration & Interoperability",
    "website": "meridian-labs.example",
    "hq": "Bengaluru, India",
    "headcount": 64,
    "founded": "2012",
    "currency": "USD",
    "about": ("Meridian Integration Labs is an integration and interoperability specialist that connects clinical, financial and "
              "operational systems for regulated enterprises. Our 60+ engineers run 24x7 integration operations from Bengaluru "
              "and Dubai, and have delivered more than 2,000 production interfaces across healthcare and financial services."),
}

MERIDIAN_CAPS = [
    ["Integration Services", "Healthcare Interoperability", "Design, build and support HL7 v2.x and FHIR interfaces between HIS, LIS, RIS, PACS and EHR systems, including ADT, ORM, ORU, SIU and DFT message flows.", "Rhapsody, Mirth Connect, HL7 v2, HL7 FHIR, MLLP, IHE", "Healthcare", "Core"],
    ["Integration Services", "Integration Platform Engineering", "Engineering on enterprise integration engines: routes, filters, mappers, transformation logic, custom scripting and reusable integration patterns.", "Rhapsody, MuleSoft, Boomi, TIBCO, Kafka, REST APIs, SOAP", "Healthcare; BFSI; Logistics", "Core"],
    ["Managed Services", "Integration Managed Services (24x7)", "Outcome-based run services for integration platforms: proactive monitoring, queue and alert management, incident and problem management with RCA, SLA reporting.", "ServiceNow, Splunk, Prometheus/Grafana, ITIL", "Healthcare; BFSI", "Core"],
    ["Managed Services", "Platform Lifecycle Management", "Version upgrades, patching, performance tuning, capacity planning and high-availability configuration for integration engines.", "Rhapsody, MuleSoft, Linux, Windows Server, VMware", "Healthcare; BFSI", "Core"],
    ["API & Digital", "API Management & Gateways", "API strategy, design-first API development, gateway configuration, developer portals and API security with OAuth 2.0.", "Apigee, MuleSoft, Kong, OAuth 2.0, OpenID Connect", "BFSI; Retail; Government", "Established"],
    ["API & Digital", "Event-driven Integration", "Event streaming and asynchronous integration patterns using Kafka and message brokers for real-time data exchange.", "Kafka, RabbitMQ, Azure Integration Services", "BFSI; Logistics", "Established"],
    ["Quality Engineering", "Interface Testing & Validation", "Interface validation, unit, integration and end-to-end testing, test evidence and go-live readiness reporting; automated regression packs for HL7 interfaces.", "HL7 v2, HL7 FHIR, Postman, Jenkins", "Healthcare", "Established"],
    ["Transition Services", "Service Transition & Knowledge Transfer", "Structured transition from incumbent vendors: knowledge acquisition, interface baselining, runbook validation and governance setup within 4–6 weeks.", "ITIL, Confluence", "Healthcare; BFSI; Government", "Core"],
    ["Advisory", "Integration Architecture Assessment", "Assess integration landscapes, rationalise interfaces and define target-state integration architecture and roadmaps.", "TOGAF, ArchiMate, Enterprise Architecture", "Healthcare; Government", "Established"],
    ["Security & Compliance", "Secure Integration & Compliance", "Secure transport (TLS, certificate management), audit logging, access control and alignment with healthcare data-privacy regulations and HIPAA principles.", "PKI/TLS, IAM, SIEM", "Healthcare; BFSI", "Established"],
    ["Security & Compliance", "Business Continuity for Integration", "Disaster recovery design, failover testing, backup validation and BCP drills for integration platforms.", "VMware, Azure, AWS", "Healthcare; BFSI", "Established"],
    ["DevOps", "Integration DevOps & Release Management", "Version control, CI/CD pipelines for interfaces, change and release governance with CAB participation and rollback capability.", "Git, CI/CD, Jenkins, Azure DevOps", "Healthcare; BFSI", "Established"],
    ["Data", "Clinical Data Mapping & Migration", "Data-mapping specifications, MRN conversions and data migration for healthcare application upgrades.", "SQL Server, Oracle DB, HL7 FHIR", "Healthcare", "Emerging"],
]

ROLE_MIX_M = [
    ("Engagement Manager", "L5", 2), ("Project Manager", "L4", 3), ("Integration Architect", "L5", 4), ("Solution Architect", "L5", 2),
    ("Senior Integration Engineer", "L4", 12), ("Integration Engineer", "L3", 16), ("Support Engineer (L2)", "L3", 9),
    ("QA Engineer", "L3", 6), ("DevOps Engineer", "L3", 4), ("Business Analyst", "L3", 3), ("Security Specialist", "L4", 2),
    ("Technical Writer", "L2", 1),
]
SKILLSETS_M = {
    "Engagement Manager": ["ITIL", "Agile/Scrum", "Stakeholder Management", "Managed Services", "ServiceNow"],
    "Project Manager": ["Agile/Scrum", "PMP", "Jira", "ITIL", "Risk Management"],
    "Integration Architect": ["Rhapsody", "HL7 v2", "HL7 FHIR", "MuleSoft", "Enterprise Architecture", "TOGAF", "Kafka"],
    "Solution Architect": ["Microsoft Azure", "AWS", "Enterprise Architecture", "TOGAF", "REST APIs", "Apigee"],
    "Senior Integration Engineer": ["Rhapsody", "HL7 v2", "HL7 FHIR", "MLLP", "JavaScript/TypeScript", "SQL Server", "Mirth Connect", "MuleSoft", "Boomi"],
    "Integration Engineer": ["HL7 v2", "HL7 FHIR", "Mirth Connect", "MuleSoft", "Boomi", "REST APIs", "SOAP", "JSON", "XML", "Java"],
    "Support Engineer (L2)": ["Rhapsody", "Splunk", "ServiceNow", "Linux", "HL7 v2", "Incident Management", "Prometheus/Grafana"],
    "QA Engineer": ["Interface Testing", "Postman", "HL7 v2", "HL7 FHIR", "Test Automation", "Jenkins"],
    "DevOps Engineer": ["CI/CD", "Git", "Kubernetes", "Docker", "Microsoft Azure", "Terraform", "Linux"],
    "Business Analyst": ["Requirements Analysis", "Data Mapping", "HL7 v2", "Clinical Workflows", "Confluence"],
    "Security Specialist": ["PKI/TLS", "IAM", "SIEM", "ISO 27001", "HIPAA"],
    "Technical Writer": ["Technical Documentation", "Runbooks", "Confluence"],
}


def build_meridian(out: Path, rnd: random.Random):
    _xlsx(out / "01_capabilities.xlsx", "Capabilities",
          ["Service Line", "Capability", "Description", "Technologies", "Industries", "Maturity"], MERIDIAN_CAPS, [20, 32, 70, 40, 26, 12])

    rows = []
    n = 1
    for role, lvl, count in ROLE_MIX_M:
        for _ in range(count):
            pool = SKILLSETS_M[role]
            k = min(len(pool), rnd.randint(4, 6))
            skills = rnd.sample(pool, k)
            # Rhapsody depth is the differentiator: keep it present but not universal
            if role in ("Senior Integration Engineer", "Integration Engineer") and "Rhapsody" not in skills and rnd.random() < 0.35:
                skills.insert(0, "Rhapsody")
            skills_txt = ", ".join(f"{s} ({rnd.choice(['3', '4', '4', '5'] if lvl in ('L4', 'L5') else ['2', '3', '3', '4'])})" for s in skills)
            years = {"L2": rnd.randint(2, 4), "L3": rnd.randint(3, 7), "L4": rnd.randint(7, 12), "L5": rnd.randint(11, 19)}[lvl]
            loc = "Dubai, UAE" if rnd.random() < 0.22 else "Bengaluru, India"
            certs = []
            if "Rhapsody" in skills and rnd.random() < 0.5:
                certs.append("Rhapsody Certified Developer")
            if role in ("Engagement Manager", "Project Manager") and rnd.random() < 0.7:
                certs.append(rnd.choice(["PMP", "ITIL v4 Foundation", "PRINCE2"]))
            if role in ("Integration Architect", "Solution Architect") and rnd.random() < 0.6:
                certs.append(rnd.choice(["TOGAF 9 Certified", "MuleSoft Certified Integration Architect", "Azure Solutions Architect Expert"]))
            if role == "Security Specialist":
                certs.append("CISSP")
            avail = rnd.choice([100, 100, 50, 0, 0, 25, 100, 0])
            rows.append([f"MIL-{n:03d}", role, lvl, years, loc, "Integration Practice" if "Integration" in role or role.startswith("Senior") else "Delivery",
                         skills_txt, "; ".join(certs), f"{avail}%"])
            n += 1
    _xlsx(out / "02_skills_matrix.xlsx", "Skills Matrix",
          ["Emp ID", "Designation", "Grade", "Total Exp (yrs)", "Base Location", "Practice", "Primary Skills (rating 1-5)", "Certifications", "Available %"],
          rows, [10, 28, 8, 12, 18, 20, 70, 40, 12], title="Meridian Integration Labs – Skills Matrix (Q3 2026)")

    _xlsx(out / "03_certifications.xlsx", "Certifications",
          ["Certification", "Type", "Issued By", "No. of Certified Employees", "Valid Until"],
          [["ISO 9001:2015 Quality Management", "Company", "Bureau Veritas", "", "2027-11-30"],
           ["ISO/IEC 27001:2022 Information Security", "Company", "Bureau Veritas", "", "2027-06-30"],
           ["ISO 22301:2019 Business Continuity", "Company", "TÜV SÜD", "", "2026-12-31"],
           ["Rhapsody Certified Partner", "Partner", "Lyniate (Rhapsody)", 9, "2027-03-31"],
           ["MuleSoft Partner", "Partner", "Salesforce MuleSoft", 7, "2027-01-31"],
           ["Boomi Partner", "Partner", "Boomi", 4, "2026-12-31"],
           ["HL7 FHIR Proficiency", "Individual", "HL7 International", 6, ""],
           ["ITIL v4 Foundation", "Individual", "PeopleCert", 14, ""],
           ["TOGAF 9 Certified", "Individual", "The Open Group", 3, ""],
           ["CISSP", "Individual", "ISC2", 2, ""]],
          [44, 14, 24, 22, 14])

    _xlsx(out / "04_rate_card.xlsx", "Rate Card",
          ["Role", "Level", "Delivery Location", "Day Rate (USD)", "Day Cost (USD)"],
          [["Engagement Manager", "L5", "Onsite", 1150, 640], ["Engagement Manager", "L5", "Offshore", 560, 260],
           ["Project Manager", "L4", "Onsite", 950, 520], ["Project Manager", "L4", "Offshore", 430, 200],
           ["Solution Architect", "L5", "Onsite", 1200, 680], ["Solution Architect", "L5", "Offshore", 600, 290],
           ["Integration Engineer", "L4", "Onsite", 820, 450], ["Integration Engineer", "L4", "Offshore", 380, 175],
           ["Integration Engineer", "L3", "Offshore", 300, 135], ["Support Engineer", "L3", "Offshore", 260, 115],
           ["Support Engineer", "L3", "Onsite", 640, 360], ["QA Engineer", "L3", "Offshore", 250, 110],
           ["DevOps Engineer", "L3", "Offshore", 330, 150], ["Business Analyst", "L3", "Onsite", 720, 400],
           ["Business Analyst", "L3", "Offshore", 300, 135], ["Security Specialist", "L4", "Offshore", 420, 200],
           ["Technical Writer", "L2", "Offshore", 180, 80]],
          [26, 8, 18, 16, 16])

    _xlsx(out / "05_case_studies.xlsx", "Case Studies",
          ["Project Name", "Client", "Industry", "Region", "Year", "Summary", "Technologies", "Outcomes"],
          [["24x7 Integration Managed Services for a 12-hospital network", "Regional hospital group (GCC)", "Healthcare", "UAE", "2024",
            "Took over run and change services for 340 Rhapsody interfaces from an incumbent in five weeks; outcome-based SLAs with P1 response under 15 minutes.",
            "Rhapsody, HL7 v2, HL7 FHIR, ServiceNow, Splunk", "99.95% interface availability; P1 MTTR down 48%; 120 new interfaces delivered in year one"],
           ["FHIR interoperability layer for a national health exchange", "Government health agency", "Healthcare", "Middle East", "2023",
            "Designed and built FHIR R4 APIs and HL7 v2-to-FHIR transformation for EHR interoperability across public and private providers.",
            "HL7 FHIR, HL7 v2, Rhapsody, Apigee, OAuth 2.0", "38 facilities connected; 2.1M messages/day; onboarding time per facility cut from 10 to 3 weeks"],
           ["Rhapsody upgrade and HA re-platforming", "Tertiary care hospital", "Healthcare", "India", "2025",
            "Upgraded Rhapsody across DEV/UAT/PROD, introduced clustering and failover, and automated regression testing for 180 interfaces.",
            "Rhapsody, Linux, VMware, Jenkins", "Zero-downtime cut-over; failover tested quarterly; regression cycle from 9 days to 1"],
           ["Core banking API platform", "Mid-size retail bank", "BFSI", "Kuwait", "2024",
            "Built an API management layer exposing core banking services to mobile and partner channels with OAuth 2.0 security.",
            "MuleSoft, Apigee, OAuth 2.0, Kafka", "65 APIs live; partner onboarding 4x faster"],
           ["Interface rationalisation assessment", "Healthcare payer", "Healthcare", "UAE", "2025",
            "Assessed 420 point-to-point interfaces and defined a target integration architecture with phased roadmap.",
            "TOGAF, Enterprise Architecture, MuleSoft", "31% of interfaces decommissioned; three-year roadmap adopted"],
           ["Lab and radiology integration (LIS/RIS/PACS)", "Diagnostic chain", "Healthcare", "India", "2022",
            "Integrated LIS, RIS and PACS with HIS using ORM/ORU messaging and DICOM work-lists.",
            "HL7 v2, DICOM, Mirth Connect", "Turnaround time for results reduced 35%"]],
          [40, 26, 12, 12, 8, 70, 36, 50])

    _xlsx(out / "06_accelerators.xlsx", "Accelerators",
          ["Accelerator", "Description", "Technologies", "Effort Saving %"],
          [["HL7 Interface Factory", "Template library of 60+ HL7 v2/FHIR routes, mappers and test packs for common ADT/ORM/ORU/SIU/DFT flows.", "Rhapsody, HL7 v2, HL7 FHIR", 25],
           ["Integration Command Centre", "Pre-built dashboards, alert rules and runbooks for 24x7 monitoring of integration engines and queues.", "Splunk, Prometheus/Grafana, ServiceNow", 20],
           ["Transition-in-a-Box", "Checklists, interface baselining scripts and KT templates that compress vendor transition to 4 weeks.", "ITIL", 30],
           ["FHIR Starter Kit", "Reusable FHIR R4 profiles, terminology mappings and API security policies.", "HL7 FHIR, Apigee, OAuth 2.0", 20]],
          [28, 80, 36, 14])

    # A narrative case study as a Word document (tests document parsing)
    d = Document()
    d.add_heading("Case study: Outcome-based integration support for a public health authority", level=1)
    d.add_paragraph("Client: Public health regulator (GCC)")
    d.add_paragraph("Industry: Healthcare – Public sector")
    d.add_heading("Challenge", level=2)
    d.add_paragraph("The authority relied on a single named contractor to keep its Rhapsody integration engine running. Incidents were "
                    "resolved slowly and new interface requests waited months, putting clinical data exchange at risk.")
    d.add_heading("What we did", level=2)
    for t in ["Transitioned 210 HL7 v2 and FHIR interfaces in four weeks using our Transition-in-a-Box accelerator.",
              "Introduced 24x7 monitoring with queue and alert dashboards and ITIL-aligned incident, problem and change management.",
              "Delivered new interfaces through two-week sprints governed by the client's change advisory board."]:
        d.add_paragraph(t, style="List Bullet")
    d.add_heading("Results", level=2)
    for t in ["P1 incidents resolved within 2 hours in 98% of cases.", "Interface backlog cleared in 5 months.",
              "Successful DR failover drill every quarter."]:
        d.add_paragraph(t, style="List Bullet")
    d.save(out / "07_case_study_public_health.docx")

    build_template(out / "08_proposal_template.docx", MERIDIAN["name"], RGBColor(0x0E, 0x5A, 0x8A), "Linking systems. Protecting care.")


# ============================================================================ Cobalt Ridge
COBALT = {
    "name": "Cobalt Ridge Technologies",
    "slug": "cobalt-ridge",
    "industry": "IT Services – Cloud, Data & Digital Banking",
    "website": "cobaltridge.example",
    "hq": "Pune, India",
    "headcount": 72,
    "founded": "2015",
    "currency": "USD",
    "about": ("Cobalt Ridge Technologies builds cloud-native digital banking, data and AI solutions for banks and insurers in the "
              "Middle East and India, from mobile channels to fraud analytics."),
}
COBALT_CAPS = [
    ["Digital Channels", "Mobile & Internet Banking Engineering", "Native iOS/Android and web banking apps, SDK integration and secure mobile development.", "iOS, Android, React, Java", "BFSI", "Core"],
    ["Data & AI", "Fraud & Risk Analytics", "Transaction-monitoring models, rule engines and investigator dashboards for card and digital-channel fraud.", "Python, Machine Learning, Databricks, Power BI", "BFSI", "Established"],
    ["Data & AI", "Graph & Link Analytics", "Entity resolution and graph analytics to expose mule networks and linked accounts.", "Graph Analytics, Python, Neo4j", "BFSI", "Emerging"],
    ["Cloud", "Cloud Migration & Landing Zones", "Secure landing zones, migration factories and managed cloud operations on Azure and AWS.", "Microsoft Azure, AWS, Terraform, Kubernetes", "BFSI; Retail", "Core"],
    ["Cloud", "Managed Cloud Operations", "24x7 cloud operations, monitoring, patching, backup and DR for regulated workloads.", "Microsoft Azure, AWS, Dynatrace, ServiceNow", "BFSI", "Established"],
    ["Security", "Identity & Access Management", "IAM, MFA and step-up authentication integration for digital channels.", "IAM, MFA, OAuth 2.0, OpenID Connect", "BFSI", "Established"],
    ["Security", "Application Security & DevSecOps", "Secure SDLC, code scanning, penetration testing coordination and mobile app hardening.", "CI/CD, Mobile App Security SDK, SIEM", "BFSI", "Established"],
    ["Integration", "API Integration", "REST API integration between channels, core banking and third-party services.", "REST APIs, Kafka, MuleSoft", "BFSI", "Established"],
    ["Data & AI", "Regulatory Reporting & BI", "Regulatory and management reporting, data warehousing and KPI dashboards.", "Power BI, Snowflake, SQL Server", "BFSI; Insurance", "Core"],
]
SKILL_COLS_C = ["Java", "Python", "React", "iOS", "Android", "Microsoft Azure", "AWS", "Kubernetes", "Terraform", "Machine Learning",
                "Databricks", "Power BI", "Graph Analytics", "IAM", "MFA", "Kafka", "REST APIs", "Mobile App Security SDK", "SIEM", "Fraud Monitoring Platform"]
ROLE_MIX_C = [
    ("Delivery Manager", "Principal", 3), ("Scrum Master", "Senior", 3), ("Cloud Architect", "Principal", 3), ("Data Scientist", "Senior", 6),
    ("Mobile Developer", "Senior", 9), ("Full Stack Developer", "Mid", 14), ("Data Engineer", "Mid", 8), ("Cloud Engineer", "Mid", 8),
    ("Security Engineer", "Senior", 4), ("QA Engineer", "Mid", 8), ("Business Analyst", "Senior", 4), ("Support Engineer", "Mid", 2),
]
FOCUS_C = {
    "Delivery Manager": ["Java", "REST APIs"], "Scrum Master": ["Java"], "Cloud Architect": ["Microsoft Azure", "AWS", "Kubernetes", "Terraform"],
    "Data Scientist": ["Python", "Machine Learning", "Databricks", "Graph Analytics", "Fraud Monitoring Platform"],
    "Mobile Developer": ["iOS", "Android", "Mobile App Security SDK", "REST APIs"], "Full Stack Developer": ["Java", "React", "REST APIs", "Kafka"],
    "Data Engineer": ["Python", "Databricks", "Power BI", "Kafka"], "Cloud Engineer": ["Microsoft Azure", "AWS", "Kubernetes", "Terraform"],
    "Security Engineer": ["IAM", "MFA", "SIEM", "Mobile App Security SDK"], "QA Engineer": ["Java", "REST APIs"],
    "Business Analyst": ["Power BI", "Fraud Monitoring Platform"], "Support Engineer": ["Microsoft Azure", "REST APIs"],
}


def build_cobalt(out: Path, rnd: random.Random):
    _xlsx(out / "capabilities.xlsx", "Offerings", ["Practice", "Offering", "What we do", "Tools & Platforms", "Verticals", "Strength"],
          COBALT_CAPS, [18, 34, 70, 40, 16, 12])
    header = ["Employee Code", "Job Title", "Band", "Experience", "Location"] + SKILL_COLS_C + ["Certs", "Allocation"]
    rows = []
    n = 1
    for role, band, count in ROLE_MIX_C:
        for _ in range(count):
            ratings = []
            for s in SKILL_COLS_C:
                if s in FOCUS_C[role]:
                    ratings.append(rnd.choice([3, 4, 4, 5]))
                elif rnd.random() < 0.12:
                    ratings.append(rnd.choice([1, 2, 3]))
                else:
                    ratings.append("")
            exp = {"Principal": rnd.randint(12, 18), "Senior": rnd.randint(7, 11), "Mid": rnd.randint(3, 6)}[band]
            loc = "Kuwait City" if rnd.random() < 0.12 else "Pune"
            certs = []
            if role == "Cloud Architect":
                certs.append("Azure Solutions Architect Expert")
            if role == "Security Engineer" and rnd.random() < 0.6:
                certs.append("CEH")
            if role == "Data Scientist" and rnd.random() < 0.4:
                certs.append("Databricks ML Associate")
            alloc = rnd.choice(["Bench", "50%", "100%", "100%", "Bench", "75%"])
            rows.append([f"CRT{1000 + n}", role, band, f"{exp} yrs", loc] + ratings + [", ".join(certs), alloc])
            n += 1
    _xlsx(out / "people_skills_grid.xlsx", "Skills Grid", header, rows, [12, 22, 10, 10, 12] + [9] * len(SKILL_COLS_C) + [30, 10])
    _xlsx(out / "certifications.xlsx", "Certs", ["Name", "Category", "Body", "Holders"],
          [["ISO/IEC 27001:2013", "Company", "BSI", ""], ["Microsoft Solutions Partner – Azure Infrastructure", "Partner", "Microsoft", 18],
           ["AWS Select Tier Partner", "Partner", "AWS", 9], ["Databricks Consulting Partner", "Partner", "Databricks", 5],
           ["Azure Solutions Architect Expert", "Individual", "Microsoft", 3], ["CEH", "Individual", "EC-Council", 2]], [48, 14, 18, 10])
    _xlsx(out / "rate_card_hourly.xlsx", "Rates", ["Resource", "Band", "Shore", "Hourly Rate (USD)", "Hourly Cost (USD)"],
          [["Delivery Manager", "Principal", "Onsite", 135, 74], ["Delivery Manager", "Principal", "Offshore", 62, 30],
           ["Cloud Architect", "Principal", "Offshore", 70, 34], ["Cloud Architect", "Principal", "Onsite", 150, 82],
           ["Data Scientist", "Senior", "Offshore", 58, 27], ["Mobile Developer", "Senior", "Offshore", 48, 22],
           ["Full Stack Developer", "Mid", "Offshore", 38, 17], ["Data Engineer", "Mid", "Offshore", 40, 18],
           ["Cloud Engineer", "Mid", "Offshore", 40, 18], ["Security Engineer", "Senior", "Offshore", 55, 26],
           ["QA Engineer", "Mid", "Offshore", 30, 13], ["Business Analyst", "Senior", "Onsite", 95, 52],
           ["Scrum Master", "Senior", "Offshore", 45, 21], ["Support Engineer", "Mid", "Offshore", 32, 14]], [24, 12, 10, 18, 18])
    _xlsx(out / "case_studies.xlsx", "References", ["Title", "Customer", "Sector", "Country", "Year", "Description", "Technology", "Results"],
          [["Card fraud analytics platform", "Retail bank", "BFSI", "Kuwait", "2025", "Built ML-based card-not-present fraud scoring with an investigator dashboard and rule engine.",
            "Python, Machine Learning, Databricks, Power BI", "Fraud losses down 27%; false positives down 18%"],
           ["Mobile banking app rebuild", "Islamic bank", "BFSI", "UAE", "2024", "Rebuilt iOS and Android banking apps with device binding and MFA step-up.",
            "iOS, Android, MFA, IAM", "4.7 app-store rating; login failures halved"],
           ["Azure landing zone for a payments company", "Payments processor", "BFSI", "India", "2023", "Regulated landing zone, DR design and managed cloud operations.",
            "Microsoft Azure, Terraform, Kubernetes", "RTO 1 hour; 99.95% uptime"],
           ["Mule-account detection pilot", "Digital bank", "BFSI", "India", "2025", "Graph analytics pilot to link devices, accounts and beneficiaries.",
            "Graph Analytics, Python", "Identified 3 mule rings in 6 weeks"],
           ["Regulatory reporting modernisation", "Insurer", "Insurance", "UAE", "2022", "Data warehouse and BI for regulatory returns.",
            "Snowflake, Power BI", "Report preparation from 10 days to 2"]], [34, 18, 10, 10, 8, 60, 34, 40])
    build_template(out / "proposal_template.docx", COBALT["name"], RGBColor(0x22, 0x3A, 0x8F), "Engineering trust in digital banking.")


# ============================================================================ template
def _field(paragraph, instr: str):
    r = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    r._r.append(f1); r._r.append(it); r._r.append(f2)


def build_template(path: Path, company: str, color: RGBColor, tagline: str):
    d = Document()
    st = d.styles
    st["Normal"].font.name = "Calibri"
    st["Normal"].font.size = Pt(10.5)
    for name, size in (("Heading 1", 18), ("Heading 2", 13), ("Heading 3", 11)):
        st[name].font.name = "Calibri"
        st[name].font.size = Pt(size)
        st[name].font.color.rgb = color
        st[name].font.bold = True
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    hp = sec.header.paragraphs[0]
    hp.text = f"{company}  |  Confidential proposal"
    hp.runs[0].font.size = Pt(8)
    hp.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run(f"{company} — Page ")
    r.font.size = Pt(8)
    _field(fp, "PAGE")
    # cover
    for _ in range(6):
        d.add_paragraph()
    p = d.add_paragraph()
    r = p.add_run(company.upper())
    r.bold = True; r.font.size = Pt(12); r.font.color.rgb = color
    p = d.add_paragraph()
    r = p.add_run("{{TITLE}}")
    r.bold = True; r.font.size = Pt(26)
    p = d.add_paragraph()
    r = p.add_run("Proposal submitted to {{CLIENT}}")
    r.font.size = Pt(13)
    p = d.add_paragraph()
    r = p.add_run("{{DATE}}")
    r.font.size = Pt(11); r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    d.add_paragraph()
    p = d.add_paragraph()
    r = p.add_run(tagline)
    r.italic = True; r.font.color.rgb = color
    d.add_page_break()
    d.add_paragraph("{{CONTENT}}")
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)


def main():
    rnd = random.Random(42)
    build_meridian(ROOT / "meridian-integration-labs", rnd)
    build_cobalt(ROOT / "cobalt-ridge-technologies", random.Random(7))
    (ROOT / "README.md").write_text(
        "# Sample company packs\n\nTwo fictional companies used to demo RevenueForge. Upload these files during company onboarding,\n"
        "as a test of the onboarding flow.\n\n"
        "- **Meridian Integration Labs**: integration & healthcare interoperability specialist (list-style skills matrix, day rates).\n"
        "- **Cobalt Ridge Technologies**: cloud, data & digital-banking company (rating-grid skills matrix, hourly rates).\n\n"
        "All names, people and projects are fictional.\n", encoding="utf-8")
    print("written to", ROOT)


if __name__ == "__main__":
    main()
