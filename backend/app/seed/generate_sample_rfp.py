"""Builds a realistic but entirely fictional RFP for testing RevenueForge without any client's confidential document.

Issuer, people, figures and systems are invented. Run:  python -m app.seed.generate_sample_rfp  (from backend/)
Writes sample-data/sample-rfps/harborline-integration-platform-rfp.docx
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ISSUER = "Harborline Regional Transit Authority"
SHORT = "HRTA"
TITLE = "Integration Platform Modernisation and Managed Services"
REF = "HRTA-RFP-2026-041"


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    doc.add_paragraph()
    return t


def reqs(doc, prefix, items):
    table(doc, ["ID", "Requirement", "Priority"], [(f"{prefix}-{i + 1:02d}", t, p) for i, (t, p) in enumerate(items)])


def build(out: Path):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(ISSUER)
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    doc.add_paragraph("Request for Proposal", style="Title").alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(TITLE)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(14)
    p = doc.add_paragraph(f"Reference: {REF}\nIssue date: 1 September 2026\nProposal due date: 30 November 2026, 14:00 local time")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("This RFP is a fictional document created for product testing. The issuer, people, systems and figures are invented.").italic = True
    doc.add_page_break()

    doc.add_heading("1. Introduction", 1)
    doc.add_paragraph(
        f"{ISSUER} ({SHORT}) operates bus, light rail and ferry services for a metropolitan region of 2.4 million residents. "
        f"{SHORT} carries 310,000 passenger journeys on a typical weekday across 96 routes and 4 ferry terminals."
    )
    doc.add_paragraph(
        f"{SHORT} invites proposals from qualified bidders to modernise its integration platform and to operate it as a managed "
        "service for an initial term of three years, with two optional one-year extensions."
    )
    doc.add_heading("1.1 Background", 2)
    doc.add_paragraph(
        "Today, 41 business applications exchange data through a mix of point-to-point file transfers, a legacy enterprise service "
        "bus that reaches end of support in June 2027, and scripts maintained by individual teams. Fare collection, vehicle "
        "telemetry, the passenger information system and the finance platform are the most critical consumers."
    )
    doc.add_paragraph(
        "Incidents caused by failed or delayed interfaces averaged 23 per month in the last financial year. Real-time vehicle "
        "positions reach the passenger app with a median delay of 95 seconds, which drives complaints."
    )
    doc.add_heading("1.2 Objectives", 2)
    for t in [
        "Replace the legacy service bus and point-to-point interfaces with a governed, API-led and event-driven integration platform.",
        "Deliver vehicle positions and service alerts to passenger channels within 5 seconds end to end.",
        "Reduce interface-related incidents by 70% within 12 months of go-live.",
        "Give business teams self-service access to approved APIs through a developer portal.",
        "Transfer day-to-day operation to a managed service with clear service levels and continuous improvement.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("2. Scope of Work", 1)
    doc.add_paragraph("The scope covers four workstreams. Bidders shall address every workstream in their response.")
    table(doc, ["Workstream", "Summary"], [
        ("A. Platform build", "Design, build and harden the target integration platform across two cloud regions."),
        ("B. Interface migration", "Migrate the 41 in-scope applications and 186 interfaces from the legacy bus and scripts."),
        ("C. API and event enablement", "Publish reusable APIs and real-time event streams with a developer portal."),
        ("D. Managed services", "Operate, support and continuously improve the platform for three years."),
    ])

    doc.add_heading("3. Functional Requirements", 1)
    doc.add_heading("3.1 Integration and APIs", 2)
    reqs(doc, "INT", [
        ("The platform shall support REST, SOAP, SFTP, JMS and Kafka-compatible interfaces.", "Mandatory"),
        ("The bidder shall provide an API gateway with rate limiting, OAuth 2.0 and mutual TLS.", "Mandatory"),
        ("The solution shall provide a developer portal where internal teams can discover, test and subscribe to APIs.", "Mandatory"),
        ("The bidder shall migrate all 186 existing interfaces with no loss of data or functionality.", "Mandatory"),
        ("The platform must support canonical data models for vehicles, trips, stops, fares and customers.", "Mandatory"),
        ("The solution should provide low-code mapping tools for business analysts.", "Desirable"),
        ("The bidder shall integrate the fare collection system (FareFlow 7) and the finance platform (Oracle Fusion).", "Mandatory"),
        ("The platform shall support versioning and deprecation policies for APIs.", "Mandatory"),
        ("The solution should support GraphQL for the passenger app backend.", "Desirable"),
        ("The bidder shall deliver reusable connectors for SAP SuccessFactors and ServiceNow.", "Mandatory"),
    ])
    doc.add_heading("3.2 Real-time events", 2)
    reqs(doc, "EVT", [
        ("The platform shall ingest vehicle telemetry from 1,450 vehicles at one message per vehicle every 2 seconds.", "Mandatory"),
        ("Vehicle positions and service alerts must reach passenger channels within 5 seconds end to end.", "Mandatory"),
        ("The solution shall publish GTFS-Realtime feeds for trip updates, vehicle positions and alerts.", "Mandatory"),
        ("The platform shall support event replay for at least 7 days.", "Mandatory"),
        ("The bidder should provide stream processing for headway and bunching detection.", "Desirable"),
        ("The solution shall guarantee at-least-once delivery with idempotent consumers.", "Mandatory"),
    ])
    doc.add_heading("3.3 Data and analytics", 2)
    reqs(doc, "DAT", [
        ("The bidder shall land integration events in the existing Azure data lake within 15 minutes.", "Mandatory"),
        ("The platform shall maintain data lineage from source system to report.", "Desirable"),
        ("The solution shall provide dashboards for interface volumes, latency and errors.", "Mandatory"),
        ("The bidder should propose a master data approach for stops and routes.", "Desirable"),
    ])

    doc.add_heading("4. Non-Functional Requirements", 1)
    doc.add_paragraph("The solution shall meet the following service levels from the start of the managed service.")
    table(doc, ["Measure", "Target"], [
        ("Platform availability (monthly)", "99.95%"),
        ("Real-time event latency, 95th percentile", "≤ 2 seconds within the platform"),
        ("API response time, 95th percentile", "≤ 300 ms"),
        ("Recovery time objective (RTO)", "1 hour"),
        ("Recovery point objective (RPO)", "5 minutes"),
        ("Peak throughput", "4,000 messages per second sustained for 1 hour"),
    ])
    reqs(doc, "NFR", [
        ("The platform shall run active-active across two cloud regions.", "Mandatory"),
        ("The solution shall scale automatically to handle event days with three times normal volume.", "Mandatory"),
        ("All components shall be deployed through infrastructure as code and CI/CD pipelines.", "Mandatory"),
        ("The bidder shall perform load testing at 150% of peak throughput before go-live.", "Mandatory"),
        ("The solution shall support blue-green or canary releases with no planned downtime.", "Desirable"),
    ])

    doc.add_heading("5. Security and Compliance", 1)
    reqs(doc, "SEC", [
        ("The bidder shall hold a current ISO/IEC 27001 certification.", "Mandatory"),
        ("The bidder shall hold SOC 2 Type II attestation for its managed service operations.", "Mandatory"),
        ("All data shall be encrypted in transit with TLS 1.2 or higher and at rest with AES-256.", "Mandatory"),
        ("Personal data shall remain within the national jurisdiction at all times.", "Mandatory"),
        ("The solution shall integrate with HRTA's Microsoft Entra ID for single sign-on and role-based access.", "Mandatory"),
        ("The bidder shall support annual penetration testing by an independent party and remediate critical findings within 14 days.", "Mandatory"),
        ("Payment card data flows shall remain compliant with PCI DSS v4.0.", "Mandatory"),
        ("The bidder shall provide a security operations contact available 24x7.", "Mandatory"),
    ])

    doc.add_heading("6. Managed Services", 1)
    doc.add_paragraph("From go-live, the bidder shall operate the platform as a managed service under the following terms.")
    table(doc, ["Priority", "Definition", "Response", "Resolution"], [
        ("P1", "Platform down or real-time feeds unavailable", "15 minutes", "4 hours"),
        ("P2", "Major interface failed, workaround unavailable", "30 minutes", "8 hours"),
        ("P3", "Degraded performance or single interface failure with workaround", "4 hours", "3 business days"),
        ("P4", "Request or minor issue", "1 business day", "10 business days"),
    ])
    reqs(doc, "MS", [
        ("The bidder shall provide 24x7 monitoring and incident management for the platform.", "Mandatory"),
        ("The bidder shall deliver a monthly service report covering SLAs, incidents, capacity and improvements.", "Mandatory"),
        ("The bidder shall provide a named service delivery manager.", "Mandatory"),
        ("Service credits shall apply for missed availability and resolution targets.", "Mandatory"),
        ("The bidder should propose at least four continuous improvement initiatives per year.", "Desirable"),
        ("The bidder shall support up to 20 minor enhancements per quarter within the managed service fee.", "Mandatory"),
        ("The bidder shall provide a transition-out plan and knowledge transfer at the end of the contract.", "Mandatory"),
    ])

    doc.add_heading("7. Project Delivery", 1)
    doc.add_paragraph("HRTA expects the platform build and migration to complete within 12 months of contract signature.")
    table(doc, ["Milestone", "Target"], [
        ("Mobilisation and discovery complete", "Month 1"),
        ("Platform foundation live in both regions", "Month 4"),
        ("Real-time feeds live for passenger channels", "Month 6"),
        ("Wave 1 migration (fares, finance) complete", "Month 8"),
        ("All interfaces migrated, legacy bus switched off", "Month 12"),
    ])
    reqs(doc, "DEL", [
        ("The bidder shall use an agile delivery method with two-week sprints and HRTA participation in sprint reviews.", "Mandatory"),
        ("The bidder shall provide a dedicated project manager and solution architect for the whole build.", "Mandatory"),
        ("The bidder shall train at least 25 HRTA staff on the platform and the developer portal.", "Mandatory"),
        ("The bidder should provide documentation in HRTA's Confluence space.", "Desirable"),
    ])

    doc.add_heading("8. Deliverables", 1)
    table(doc, ["Deliverable", "Due"], [
        ("Solution architecture document", "Month 2"),
        ("Integration catalogue and migration plan", "Month 2"),
        ("Security design and threat model", "Month 3"),
        ("Test strategy and test reports", "Each wave"),
        ("Operations runbooks", "Before go-live"),
        ("Service management plan", "Before go-live"),
    ])

    doc.add_heading("9. Commercial Terms", 1)
    for t in [
        "Bidders shall price the build as fixed price by milestone and the managed service as a fixed monthly fee.",
        "Prices shall be quoted in Australian dollars (AUD), exclusive of GST.",
        "Payment will be made against accepted milestones, net 30 days from invoice.",
        "A performance bond of 5% of the build price is required until final acceptance.",
        "Rates for additional work shall be held for the first two years of the contract.",
        "The bidder shall carry professional indemnity insurance of at least AUD 10 million.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("10. Evaluation Criteria", 1)
    table(doc, ["Criterion", "Weight"], [
        ("Technical solution and architecture", "30%"),
        ("Delivery approach, team and experience", "20%"),
        ("Managed service model and service levels", "20%"),
        ("Price", "25%"),
        ("Social value and local presence", "5%"),
    ])

    doc.add_heading("11. Submission Instructions", 1)
    for t in [
        "Proposals shall be submitted through the HRTA e-tendering portal by 30 November 2026, 14:00 local time.",
        "The technical proposal shall not exceed 60 pages, excluding appendices.",
        "Bidders shall include three references from comparable engagements within the last five years.",
        "Clarification questions may be submitted until 6 November 2026.",
        "Proposals shall remain valid for 180 days from the submission deadline.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


if __name__ == "__main__":
    root = Path(__file__).resolve().parents[3]
    print(build(root / "sample-data" / "sample-rfps" / "harborline-integration-platform-rfp.docx"))
